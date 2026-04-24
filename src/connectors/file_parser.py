import yaml
from pathlib import Path
from docx import Document


class FileParser:
    """Extracts clean text and metadata from .txt, .md, and .docx files."""

    def parse(self, file_path: str) -> dict:
        path = Path(file_path)
        ext = path.suffix.lower()
        title = path.stem.replace("_", " ").replace("-", " ")

        result = {
            "text": "",
            "title": title,
            "file_path": str(path),
            "extension": ext,
            "frontmatter": {},
            "tags": "",
            "error": None,
        }

        try:
            if ext == ".txt":
                result["text"] = self._parse_txt(path)
            elif ext == ".md":
                text, frontmatter, fm_title, fm_tags = self._parse_markdown(path)
                result["text"] = text
                result["frontmatter"] = frontmatter
                if fm_title:
                    result["title"] = fm_title
                if fm_tags:
                    result["tags"] = fm_tags
            elif ext == ".docx":
                result["text"] = self._parse_docx(path)
            else:
                raise ValueError(f"Unsupported file extension: '{ext}'")
        except ValueError:
            raise
        except Exception as e:
            result["error"] = str(e)

        return result

    def _parse_txt(self, path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            return path.read_text(encoding="latin-1")

    def _parse_markdown(self, path: Path) -> tuple:
        content = path.read_text(encoding="utf-8")
        frontmatter = {}
        fm_title = ""
        fm_tags = ""
        body = content

        if content.startswith("---"):
            parts = content.split("---", 2)
            if len(parts) >= 3:
                try:
                    frontmatter = yaml.safe_load(parts[1]) or {}
                    body = parts[2].lstrip("\n")
                    fm_title = frontmatter.get("title", "")
                    tags = frontmatter.get("tags", "")
                    if isinstance(tags, list):
                        fm_tags = ", ".join(str(t) for t in tags)
                    elif tags:
                        fm_tags = str(tags)
                except yaml.YAMLError:
                    frontmatter = {}
                    body = content

        return body, frontmatter, fm_title, fm_tags

    def _parse_docx(self, path: Path) -> str:
        doc = Document(str(path))

        paragraph_text = "\n".join(p.text for p in doc.paragraphs)

        table_lines = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    table_lines.append(row_text)
        table_text = "\n".join(table_lines)

        parts = [paragraph_text]
        if table_text:
            parts.append(table_text)
        return "\n".join(parts)


if __name__ == "__main__":
    import tempfile
    import os

    parser = FileParser()
    tmp_files = []

    # --- .txt test ---
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".txt", delete=False, encoding="utf-8",
        prefix="my_plain_note_"
    ) as f:
        f.write("This is a plain text note.\nIt has two lines.")
        txt_path = f.name
    tmp_files.append(txt_path)

    # --- .md test with frontmatter ---
    with tempfile.NamedTemporaryFile(
        mode="w", suffix=".md", delete=False, encoding="utf-8",
        prefix="my_markdown_note_"
    ) as f:
        f.write(
            "---\n"
            "title: My Great Note\n"
            "tags:\n  - python\n  - mcp\n"
            "---\n\n"
            "# Introduction\n\nThis is the body of the markdown note."
        )
        md_path = f.name
    tmp_files.append(md_path)

    # --- .docx test ---
    from docx import Document as DocxDoc
    doc = DocxDoc()
    doc.add_paragraph("First paragraph of a Word document.")
    doc.add_paragraph("Second paragraph with more content.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Header A"
    table.cell(0, 1).text = "Header B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False, prefix="my_doc_") as f:
        docx_path = f.name
    doc.save(docx_path)
    tmp_files.append(docx_path)

    # --- Run and print results ---
    for path in tmp_files:
        result = parser.parse(path)
        print(f"file:        {Path(path).name}")
        print(f"  title:       {result['title']}")
        print(f"  extension:   {result['extension']}")
        print(f"  frontmatter: {result['frontmatter']}")
        print(f"  tags:        {result['tags']!r}")
        print(f"  text[:100]:  {result['text'][:100]!r}")
        print(f"  error:       {result['error']}")
        print()

    # --- Cleanup ---
    for path in tmp_files:
        os.unlink(path)
    print("Temp files cleaned up.")
